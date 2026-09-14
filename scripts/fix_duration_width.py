"""
Fix Duration column width in whitebox test reports.
The column is too narrow causing text to wrap vertically.
Sets Duration column to 1.1 inches wide.

Run from repo root:
    python scripts/fix_duration_width.py
"""

import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS_DIR = Path("docs/testings")
EMU = 914400

# Target width for Duration column
DUR_WIDTH_IN = 1.10
DUR_WIDTH_TWIPS = int(DUR_WIDTH_IN * 1440)  # twips (used by w:w/@w:w)

WB_FILES = [
    DOCS_DIR / "Alpha_Whitebox_Test_Report.docx",
    DOCS_DIR / "Beta_Whitebox_Test_Report.docx",
]

COMPACT_MAP = {
    "Alpha_Whitebox_Test_Report.docx": "Alpha_Whitebox_Test_Report_Compact.docx",
    "Beta_Whitebox_Test_Report.docx":  "Beta_Whitebox_Test_Report_Compact.docx",
}
TNR_MAP = {
    "Alpha_Whitebox_Test_Report.docx": "Alpha_Whitebox_Test_Report_TimesNewRoman12.docx",
    "Beta_Whitebox_Test_Report.docx":  "Beta_Whitebox_Test_Report_TimesNewRoman12.docx",
}


def _is_test_table(table):
    if not table.rows:
        return False
    return table.rows[0].cells[0].text.strip().lower() in ("test id", "test case id")


def _duration_col_index(table):
    if not table.rows:
        return -1
    for i, cell in enumerate(table.rows[0].cells):
        if cell.text.strip().lower() == "duration":
            return i
    return -1


def fix_duration_width(table):
    """Fix the width of the Duration column so text doesn't wrap vertically."""
    dur_idx = _duration_col_index(table)
    if dur_idx == -1:
        return False

    # 1. Fix each cell's w:tcW width
    for row in table.rows:
        if dur_idx >= len(row.cells):
            continue
        cell = row.cells[dur_idx]
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        # Remove any existing w:tcW
        for existing in tcPr.findall(qn("w:tcW")):
            tcPr.remove(existing)

        # Add correct width
        tcW = OxmlElement("w:tcW")
        tcW.set(qn("w:w"), str(DUR_WIDTH_TWIPS))
        tcW.set(qn("w:type"), "dxa")
        tcPr.append(tcW)

        # Also remove any text direction that could cause vertical text
        for txDir in tcPr.findall(qn("w:textDirection")):
            tcPr.remove(txDir)

        # Ensure paragraph alignment and no word-wrapping issues
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after  = Pt(2)
            # Make sure runs have proper font settings
            for run in para.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8)

    # 2. Fix the table grid column (w:tblGrid / w:gridCol)
    tbl = table._tbl
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        gridCols = tblGrid.findall(qn("w:gridCol"))
        if dur_idx < len(gridCols):
            gridCols[dur_idx].set(qn("w:w"), str(DUR_WIDTH_TWIPS))

    return True


def process_file(path: Path):
    print(f"\n[{path.name}]")
    doc = Document(str(path))
    fixed = 0
    for table in doc.tables:
        if _is_test_table(table):
            if fix_duration_width(table):
                fixed += 1
    doc.save(str(path))
    print(f"  Fixed Duration column width in {fixed} tables  ({DUR_WIDTH_IN}\" = {DUR_WIDTH_TWIPS} twips)")


def sync_variants(src: Path):
    name = src.name
    for mapping in (COMPACT_MAP, TNR_MAP):
        dst_name = mapping.get(name)
        if dst_name:
            shutil.copy2(str(src), str(DOCS_DIR / dst_name))
            print(f"  Synced -> {dst_name}")


def main():
    print("=" * 60)
    print(f"  Fix Duration Column Width  (target: {DUR_WIDTH_IN}\")")
    print("=" * 60)
    for path in WB_FILES:
        process_file(path)
        sync_variants(path)
    print("\n" + "=" * 60)
    print("  Done!")
    print("=" * 60)


if __name__ == "__main__":
    main()
